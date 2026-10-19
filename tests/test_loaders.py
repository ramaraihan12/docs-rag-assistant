from pathlib import Path

import pytest
from docx import Document as new_docx

from docs_rag.loaders import UnsupportedFileType, load_file

SAMPLES = Path(__file__).resolve().parent.parent / "samples"


def test_pdf_keeps_one_document_per_page():
    docs = load_file(SAMPLES / "kabut-labs-employee-handbook.pdf")

    assert [d.metadata["page"] for d in docs] == [1, 2, 3, 4, 5, 6, 7]
    assert all(d.metadata["file_name"] == "kabut-labs-employee-handbook.pdf" for d in docs)
    assert "14 working days" in docs[2].text


def test_markdown_splits_on_headings():
    docs = load_file(SAMPLES / "kebijakan-perjalanan-dinas.md")
    by_section = {d.metadata["section"]: d.text for d in docs}

    assert "Rp 350.000 per hari" in by_section["Uang harian"]
    assert all(d.metadata["page"] is None for d in docs)


def test_markdown_ignores_hashes_inside_code_fences(tmp_path: Path):
    path = tmp_path / "notes.md"
    path.write_text("# Setup\n\n```bash\n# not a heading\npip install x\n```\n", encoding="utf-8")

    docs = load_file(path)

    assert len(docs) == 1
    assert docs[0].metadata["section"] == "Setup"


def test_docx_counts_explicit_page_breaks(tmp_path: Path):
    doc = new_docx()
    doc.add_heading("Intro", level=1)
    doc.add_paragraph("first page text")
    doc.add_page_break()
    doc.add_heading("Details", level=1)
    doc.add_paragraph("second page text")
    path = tmp_path / "two-pages.docx"
    doc.save(str(path))

    docs = load_file(path)

    assert [(d.metadata["page"], d.metadata["section"]) for d in docs] == [
        (1, "Intro"),
        (2, "Details"),
    ]
    assert "second page text" in docs[1].text


def test_docx_reads_tables():
    docs = load_file(SAMPLES / "tandon-sense-t2-manual.docx")
    specs = next(d for d in docs if d.metadata["section"] == "Specifications")

    assert specs.metadata["page"] == 2
    assert "Wi-Fi 2.4 GHz" in specs.text


def test_upload_name_overrides_temp_file_name(tmp_path: Path):
    path = tmp_path / "tmp123.txt"
    path.write_text("plain text", encoding="utf-8")

    docs = load_file(path, file_name="notes.txt")

    assert docs[0].metadata["file_name"] == "notes.txt"


def test_rejects_unknown_extension(tmp_path: Path):
    path = tmp_path / "sheet.xlsx"
    path.write_bytes(b"not really a spreadsheet")

    with pytest.raises(UnsupportedFileType):
        load_file(path)
