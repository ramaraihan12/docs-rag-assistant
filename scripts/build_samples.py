"""Rebuild the PDF and DOCX files in samples/.

Kabut Labs and the Tandon Sense T2 are made up. The text lives here so the binary
files can be regenerated and diffed in review. Each list item is one page, which
keeps the page numbers in evals/questions.jsonl stable.

    uv run python scripts/build_samples.py
"""

from pathlib import Path

from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

SAMPLES = Path(__file__).resolve().parent.parent / "samples"

# (heading, [paragraphs]) per section, list of sections per page
HANDBOOK: list[list[tuple[str, list[str]]]] = [
    [
        (
            "Welcome to Kabut Labs",
            [
                "Kabut Labs is a small hardware and software company in Salatiga, Central Java. "
                "We design low-power sensors for homes and small farms and the app that goes "
                "with them. There are 34 of us, split across firmware, cloud, mobile, support "
                "and operations.",
                "This handbook explains how we work: hours, leave, pay, equipment and the "
                "rules we hold each other to. It applies to everyone on a permanent or "
                "fixed-term contract. Interns get a shorter version from the People team.",
                "If something here is unclear or out of date, tell the People team at "
                "people@kabutlabs.example. We review the handbook every January. This is "
                "version 3.1, effective 1 February 2026.",
            ],
        ),
        (
            "Our office",
            [
                "The office is at Jl. Kemiri Raya 12, Salatiga. It is open from 07:00 to "
                "20:00 on working days. Visitors sign in at the front desk and must be "
                "accompanied in the lab area.",
            ],
        ),
    ],
    [
        (
            "Working hours",
            [
                "A full-time week is 40 hours, Monday to Friday. Core hours are 10:00 to "
                "15:00 WIB. During core hours you should be reachable on team chat and "
                "available for meetings. Outside core hours you choose when to start and "
                "finish, as long as you start between 07:00 and 10:00.",
                "Overtime is not expected. If a release or an incident means you work more "
                "than 45 hours in a week, record it in Lumbung (our HR app) and take the "
                "time back within the next month.",
            ],
        ),
        (
            "Remote work",
            [
                "After probation you may work remotely up to three days a week. Agree the "
                "days with your lead and keep them in the team calendar. Firmware and "
                "hardware roles may need to be in the lab more often during prototype builds.",
                "Working from outside Indonesia needs written approval from the Head of "
                "People at least 14 days before you travel, and is limited to 30 days per "
                "calendar year. You are responsible for a stable connection during core hours.",
            ],
        ),
    ],
    [
        (
            "Annual leave",
            [
                "Everyone gets 14 working days of paid annual leave per year. Leave builds up "
                "monthly, so you earn about 1.17 days for each month worked. New starters can "
                "use leave once they have earned it, including during probation.",
                "Request leave in Lumbung. For anything longer than two days, ask at least "
                "seven days ahead so the team can plan. Up to five unused days carry over "
                "into the next year; carried-over days expire on 31 March.",
            ],
        ),
        (
            "Sick leave",
            [
                "Tell your lead as early as you can on the day. Sick leave is paid. For "
                "absences longer than two consecutive working days we need a doctor's note, "
                "uploaded to Lumbung when you are back.",
            ],
        ),
        (
            "Parental and other leave",
            [
                "Maternity leave is three months, following Indonesian law. Fathers get ten "
                "working days of paid paternity leave, to be taken within 60 days of the "
                "birth or adoption. Bereavement leave is three working days for a close "
                "family member. Public holidays follow the government calendar.",
            ],
        ),
    ],
    [
        (
            "Pay day",
            [
                "Salaries are paid on the 25th of each month. When the 25th falls on a "
                "weekend or public holiday, we pay on the working day before it. Payslips "
                "are available in Lumbung on the same day.",
                "The religious holiday allowance (THR) is one month's salary and is paid at "
                "least 14 days before Eid al-Fitr, earlier than the legal minimum.",
            ],
        ),
        (
            "Benefits",
            [
                "We register everyone with BPJS Kesehatan and BPJS Ketenagakerjaan from the "
                "first day. On top of that, each employee has a learning budget of "
                "Rp 6,000,000 per year for courses, books, certifications and conference "
                "tickets. Unused budget does not roll over.",
                "People who work remotely at least two days a week receive an internet "
                "allowance of Rp 350,000 per month with their salary.",
            ],
        ),
    ],
    [
        (
            "Equipment",
            [
                "You get a laptop, a monitor and a headset on your first day. Equipment "
                "stays company property. When you leave, return it within five working days "
                "of your last day; the Operations team arranges a courier if you are remote.",
            ],
        ),
        (
            "IT security",
            [
                "Use the company password manager for every work account and turn on "
                "two-factor authentication wherever it is offered. Laptops must lock after "
                "five minutes of inactivity and use full-disk encryption.",
                "If your laptop or phone with work accounts is lost or stolen, report it to "
                "it@kabutlabs.example within 24 hours, even if you think it will turn up. "
                "IT will lock the device remotely and reset your sessions.",
                "Customer data, including sensor readings linked to an address, must never "
                "be copied to personal devices or personal cloud storage.",
            ],
        ),
    ],
    [
        (
            "Code of conduct",
            [
                "We expect respect in every interaction: with colleagues, customers, "
                "suppliers and the farmers we visit. Harassment, discrimination and bullying "
                "are not tolerated, in the office, online or at company events.",
                "Gifts from suppliers or partners worth more than Rp 500,000 must be "
                "declared to the People team within a week. Cash gifts of any amount are "
                "refused.",
            ],
        ),
        (
            "Raising a concern",
            [
                "Talk to your lead or the People team. If you would rather not be named, use "
                "the anonymous form linked in Lumbung; the form goes to two directors and "
                "not to your reporting line. We do not tolerate retaliation against anyone "
                "who raises a concern in good faith.",
            ],
        ),
    ],
    [
        (
            "Probation",
            [
                "New employees have a three-month probation period. Halfway through, you and "
                "your lead have a check-in to talk about how it is going. At the end, your "
                "lead confirms your contract in writing.",
            ],
        ),
        (
            "Performance and growth",
            [
                "Everyone has a one-to-one with their lead every two weeks. Formal reviews "
                "happen twice a year, in June and December, and salary changes are decided "
                "after the December review.",
            ],
        ),
        (
            "Leaving Kabut Labs",
            [
                "The notice period is 30 days after probation and seven days during "
                "probation. Your last salary, unused annual leave and any approved expense "
                "claims are paid with the next regular payroll.",
            ],
        ),
    ],
]

MANUAL: list[list[tuple[str, list[str]]]] = [
    [
        (
            "Tandon Sense T2",
            [
                "The Tandon Sense T2 measures the water level in a rooftop or ground tank "
                "with an ultrasonic sensor and sends it to the Kabut Home app over Wi-Fi. "
                "It never touches the water, so it is safe for drinking-water tanks.",
            ],
        ),
        (
            "In the box",
            [
                "Sensor unit with built-in batteries, mounting bracket, two stainless steel "
                "screws, a rubber gasket, a USB-C charging cable and a quick start card.",
            ],
        ),
        (
            "Safety",
            [
                "Do not open the housing. Do not submerge the unit; it is splash-proof, not "
                "waterproof. Charge only with a 5 V USB adapter. Keep the sensor face clean "
                "and pointed straight down at the water.",
            ],
        ),
    ],
    [
        (
            "Specifications",
            [
                "Measuring range | 25 cm to 4 m",
                "Accuracy | plus or minus 1 cm",
                "Batteries | 2 x 18650 Li-ion, 3,400 mAh each",
                "Battery life | about 9 months at one reading every 10 minutes",
                "Wireless | Wi-Fi 2.4 GHz, 802.11 b/g/n",
                "Ingress protection | IP65",
                "Operating temperature | -10 to 60 degrees Celsius",
                "Weight | 210 g",
            ],
        ),
    ],
    [
        (
            "Installation",
            [
                "Mount the T2 on the tank lid, as close to the centre as you can. Drill a "
                "32 mm hole, fit the gasket, then fix the bracket with the two screws.",
                "The sensor cannot measure anything closer than 25 cm. Leave at least 25 cm "
                "between the sensor face and the highest water level, otherwise readings "
                "near a full tank will be wrong.",
                "Do not mount it directly above the inlet pipe or the float valve. Splashing "
                "water and moving parts confuse the ultrasonic echo.",
            ],
        ),
    ],
    [
        (
            "Pairing with the app",
            [
                "Install Kabut Home from the Play Store or App Store and sign in. Press and "
                "hold the button on the side of the T2 for five seconds until the LED blinks "
                "blue. The sensor is now in pairing mode for three minutes.",
                "In the app, tap Add device, then Tandon Sense T2. Your phone joins the "
                "temporary hotspot TandonSense-XXXX, where XXXX are the last four characters "
                "of the serial number. Pick your home network and enter its password.",
                "Only 2.4 GHz networks work. If your router uses one name for both bands, "
                "turn on a separate 2.4 GHz network for the setup.",
            ],
        ),
    ],
    [
        (
            "LED status codes",
            [
                "Solid green for three seconds: a reading was taken and sent.",
                "Blinking blue: pairing mode.",
                "Blinking red: battery below 15 percent. Charge it within a week.",
                "Solid red: sensor error. Usually something is blocking the sensor face, or "
                "the water is closer than 25 cm.",
                "Alternating red and blue: firmware update in progress. Do not unplug or "
                "press the button until it stops.",
            ],
        ),
        (
            "Alerts",
            [
                "Set a low-level threshold per tank in the app, for example 20 percent. "
                "When the level drops below it you get a push notification, and a WhatsApp "
                "message if you linked a number in Settings.",
            ],
        ),
    ],
    [
        (
            "Maintenance",
            [
                "Wipe the sensor face with a dry cloth every three months. Firmware updates "
                "install over the air at 02:00 local time when the battery is above 30 percent.",
            ],
        ),
        (
            "Troubleshooting",
            [
                "Readings jump up and down: condensation on the sensor face or foam on the "
                "water. Wipe the sensor and turn on Smooth readings in the device settings.",
                "The device shows offline: check that the router is on 2.4 GHz and within "
                "range. The T2 needs a signal of at least -75 dBm at the tank.",
                "Factory reset: hold the button for 15 seconds until the LED flashes white "
                "three times. The device forgets its Wi-Fi and has to be paired again.",
            ],
        ),
    ],
    [
        (
            "Warranty",
            [
                "The T2 has an 18-month warranty from the date of purchase. The batteries are "
                "covered for 6 months. The warranty does not cover damage from submersion, a "
                "housing that has been opened, or using a charger above 5 V.",
                "To make a claim, email support@kabutlabs.example with your proof of purchase "
                "and the serial number from the label on the back. Serial numbers start with "
                "TS2-.",
            ],
        ),
    ],
]


def build_handbook(path: Path) -> None:
    styles = getSampleStyleSheet()

    def footer(canvas, doc) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.drawString(20 * mm, 12 * mm, "Kabut Labs Employee Handbook, version 3.1")
        canvas.drawRightString(190 * mm, 12 * mm, str(doc.page))
        canvas.restoreState()

    story = [Paragraph("Kabut Labs Employee Handbook", styles["Title"])]
    for number, page in enumerate(HANDBOOK):
        if number:
            story.append(PageBreak())
        for heading, paragraphs in page:
            story.append(Paragraph(heading, styles["Heading2"]))
            for text in paragraphs:
                story.append(Paragraph(text, styles["BodyText"]))
                story.append(Spacer(1, 4))

    doc = SimpleDocTemplate(
        str(path), pagesize=A4, title="Kabut Labs Employee Handbook", invariant=True
    )
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


def build_manual(path: Path) -> None:
    doc = Document()
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading("Tandon Sense T2 User Manual", level=0)

    for number, page in enumerate(MANUAL):
        if number:
            doc.add_page_break()
        for heading, paragraphs in page:
            doc.add_heading(heading, level=1)
            if heading == "Specifications":
                table = doc.add_table(rows=0, cols=2)
                table.style = "Light List"
                for row in paragraphs:
                    label, value = row.split(" | ")
                    cells = table.add_row().cells
                    cells[0].text, cells[1].text = label, value
                continue
            for text in paragraphs:
                doc.add_paragraph(text)

    doc.core_properties.title = "Tandon Sense T2 User Manual"
    doc.core_properties.author = "Kabut Labs"
    doc.save(str(path))


if __name__ == "__main__":
    build_handbook(SAMPLES / "kabut-labs-employee-handbook.pdf")
    build_manual(SAMPLES / "tandon-sense-t2-manual.docx")
    print(f"wrote samples to {SAMPLES}")
